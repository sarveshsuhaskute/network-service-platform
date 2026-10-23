#!/usr/bin/env python3
"""Generate report/CN_Project_Report.docx (BCSE308P template structure) from results/*.json and report/figures."""
import json
import os
import re

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
FIG = os.path.join(ROOT, "report", "figures")
R = json.load(open(os.path.join(ROOT, "results", "benchmarks.json")))
T = json.load(open(os.path.join(ROOT, "results", "test_results.json")))
OUT = os.path.join(ROOT, "report", "CN_Project_Report.docx")

TITLE = "Lightweight Client-Server Network Service Platform"
NAME, REG = "[Student Name]", "[Register Number]"
FONT = "Times New Roman"

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.27), Inches(11.69)             # A4
for side in ("left_margin", "right_margin"):
    setattr(sec, side, Inches(1.0))
USABLE = 8.27 - 2.0
sec.top_margin = sec.bottom_margin = Inches(1)

st = doc.styles["Normal"]
st.font.name, st.font.size = FONT, Pt(12)
st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
st.paragraph_format.space_after = Pt(6)
st.paragraph_format.line_spacing = 1.15
for name, size in (("Heading 1", 15), ("Heading 2", 13), ("Heading 3", 12)):
    h = doc.styles[name]
    h.font.name, h.font.size, h.font.bold = FONT, Pt(size), True
    h.font.color.rgb = RGBColor(0, 0, 0)
    h.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    h.element.rPr.rFonts.set(qn("w:ascii"), FONT)
    h.element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    h.paragraph_format.space_before = Pt(14 if name == "Heading 1" else 10)
    h.paragraph_format.space_after = Pt(6)
    h.paragraph_format.keep_with_next = True

counters = {"fig": 0, "tab": 0}
REF = {}


def runs(par, text, size=None, bold=None, italic=None):
    """Add text with **bold**, *italic* and `code` inline markup."""
    for tok in re.split(r"(\*\*.+?\*\*|`.+?`|\*[^*\s][^*]*?\*)", text):
        if not tok:
            continue
        if tok.startswith("**"):
            r = par.add_run(tok[2:-2]); r.bold = True
        elif tok.startswith("`"):
            r = par.add_run(tok[1:-1]); r.font.name = "Courier New"; r.font.size = Pt((size or 12) - 1.5)
            r._element.rPr.rFonts.set(qn("w:eastAsia"), "Courier New")
            continue
        elif tok.startswith("*") and len(tok) > 2:
            r = par.add_run(tok[1:-1]); r.italic = True
        else:
            r = par.add_run(tok)
        if size:
            r.font.size = Pt(size)
        if bold:
            r.bold = True
        if italic:
            r.italic = True


def P(text="", align="justify", size=None, bold=None, italic=None, after=None, indent=None, keep=False):
    p = doc.add_paragraph()
    p.alignment = {"justify": WD_ALIGN_PARAGRAPH.JUSTIFY, "center": WD_ALIGN_PARAGRAPH.CENTER,
                   "left": WD_ALIGN_PARAGRAPH.LEFT, "right": WD_ALIGN_PARAGRAPH.RIGHT}[align]
    runs(p, text, size, bold, italic)
    if after is not None:
        p.paragraph_format.space_after = Pt(after)
    if indent:
        p.paragraph_format.left_indent = Inches(indent)
    if keep:
        p.paragraph_format.keep_with_next = True
    return p


def bullets(items, style="List Bullet"):
    for it in items:
        p = doc.add_paragraph(style=style)
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.space_after = Pt(3)
        runs(p, it)


def H1(t): doc.add_heading(t, 1)
def H2(t): doc.add_heading(t, 2)
def H3(t): doc.add_heading(t, 3)


def shade(cell, hex_):
    tcPr = cell._element.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), hex_)
    tcPr.append(shd)


def caption(kind, text, label=None):
    counters[kind] += 1
    name = "Figure" if kind == "fig" else "Table"
    if label:
        REF[label] = f"{name} {counters[kind]}"
    return P(f"**{name} {counters[kind]}:** {text}", align="center", size=10.5, after=10)


def table(headers, rows, widths=None, cap=None, size=9.5, align_num=True, header_fill="D9E2F3", label=None):
    if cap:
        counters["tab"] += 1
        if label:
            REF[label] = f"Table {counters['tab']}"
        p = P(f"**Table {counters['tab']}:** {cap}", align="center", size=10.5, after=3, keep=True)
    if widths and sum(widths) > USABLE:
        widths = [w * USABLE / sum(widths) for w in widths]
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for i, h in enumerate(headers):
        c = t.rows[0].cells[i]
        c.text = ""
        runs(c.paragraphs[0], h, size=size, bold=True)
        c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        shade(c, header_fill)
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = ""
            para = cells[i].paragraphs[0]
            runs(para, str(v), size=size)
            para.paragraph_format.space_after = Pt(0)
            para.paragraph_format.line_spacing = 1.0
            if align_num and re.fullmatch(r"[-+≈~]?[\d.,]+\s*(%|ms|MB/s|s|KB|MB|x|×)?", str(v).strip()):
                para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for r_ in t.rows:
        for i, c in enumerate(r_.cells):
            c.paragraphs[0].paragraph_format.space_after = Pt(0)
            if widths:
                c.width = Inches(widths[i])
    # repeat header row on page breaks
    trPr = t.rows[0]._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader"); el.set(qn("w:val"), "true"); trPr.append(el)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)
    return t


def figure(name, width=6.0, cap=None, label=None):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_after = Pt(2)
    p.add_run().add_picture(os.path.join(FIG, name), width=Inches(min(width, USABLE)))
    if cap:
        caption("fig", cap, label or name)


def page_break():
    doc.paragraphs[-1].add_run().add_break(WD_BREAK.PAGE)


def placeholder(par, text):
    r = par.add_run(text)
    r.font.highlight_color = WD_COLOR_INDEX.YELLOW
    return r


def field(par, instr, placeholder_text="1"):
    for typ, txt in (("begin", None), (None, instr), ("separate", None), (None, placeholder_text), ("end", None)):
        r = par.add_run()
        if typ:
            f = OxmlElement("w:fldChar"); f.set(qn("w:fldCharType"), typ); r._r.append(f)
        elif txt == instr:
            it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = instr; r._r.append(it)
        else:
            r.text = txt


def f1(x): return f"{x:,.1f}"
def f2(x): return f"{x:,.2f}"
def f3(x): return f"{x:,.3f}"
def f0(x): return f"{x:,.0f}"
def pc(x): return f"{x * 100:.1f}%"


# =============================================================================== data shortcuts
E1T, E1U, E1N = R["e1"]["tcp_persistent"], R["e1"]["udp"], R["e1"]["tcp_new_conn"]
E2, E3, E3B = R["e2"], R["e3"], R["e3b"]
SF, SFC = R["e3c"]["files"], R["e3c"]["concurrent"]
DC, DCC = R["e8"]["files"], R["e8"]["concurrent"]
DS, DSC = R["e8b"]["files"], R["e8b"]["concurrent"]
E4L, E4Q = R["e4"]["latency"], R["e4"]["qps"]
E5, E6, E6B, E7 = R["e5"], R["e6"], R["e6b"], R["e7"]
E9 = R["e9"]
RTO40 = [r for r in E9["rows"] if r["loss_per_datagram"] == 0.4][0]
RTO_SPEEDUP = [r["fixed"]["mean_ms"] / r["adaptive"]["mean_ms"] for r in E9["rows"] if r["loss_per_datagram"] > 0]
M = R["machine"]


def by(rows, key, val):
    return [r for r in rows if r[key] == val][0]


peak = max(E2, key=lambda r: r["req_per_s"])
plateau = [r["req_per_s"] for r in E2 if r["clients"] in (10, 25, 50, 100)]
c100, c400 = by(E2, "clients", 100), by(E2, "clients", 400)
idle2000, idle0 = by(E6, "idle_connections", 2000), by(E6, "idle_connections", 0)
tcp1k = by(E1T, "size", 1024)
udp1k = by(E1U, "size", 1024)
udp16 = by(E1U, "size", 16)
tcp1m = by(E1T, "size", 1048576)
loss40 = by(E5, "loss_per_datagram", 0.4)
loss10 = by(E5, "loss_per_datagram", 0.1)
loss20 = by(E5, "loss_per_datagram", 0.2)


def dns_at(row, r): return [d for d in row["dns"] if d["retries"] == r][0]


dl_base100, dl_sf100, dl_dc100, dl_ds100 = (by(x, "size_mb", 100)["download_MBps"] for x in (E3, SF, DC, DS))
up100 = by(E3, "size_mb", 100)["upload_MBps"]
c8 = [by(x, "clients", 8) for x in (E3B, SFC, DCC, DSC)]
c16 = [by(x, "clients", 16) for x in (E3B, SFC, DCC, DSC)]
NQ = E5[0]['dns'][0]['queries']
MAXDEV_RAW = max(abs(r['udp_observed_loss'] - r['udp_theory_loss']) for r in E5) * 100
MAXDEV_SUCC = max(abs(d['success'] - d['theory_success']) for r in E5 for d in r['dns']) * 100
MAXSE = max((d['theory_success'] * (1 - d['theory_success']) / NQ) ** 0.5 for r in E5 for d in r['dns']) * 100
tests_pass = sum(t["status"] == "PASS" for t in T["tests"])

# =============================================================================== title page
for _ in range(2):
    P("", after=12)
P(TITLE, align="center", size=24, bold=True, after=40)
P("BCSE308P – Computer Networks", align="center", size=16, bold=True, after=30)
P("PROJECT BASED COMPONENT REPORT", align="center", size=16, bold=True, after=30)
P("by", align="center", size=13, after=14)
p = P("", align="center", size=14, after=2); placeholder(p, NAME).font.size = Pt(14)
p.runs[0].bold = True
p = P("", align="center", size=13, after=40); placeholder(p, f"({REG})")
P("School of Computer Science and Engineering", align="center", size=14, bold=True, after=70)
P("October 2026", align="center", size=14, bold=True)
page_break()

# =============================================================================== declaration
P("DECLARATION", align="center", size=14, bold=True, after=18)
p = P("", after=10)
runs(p, f"I hereby declare that the report entitled “{TITLE}” submitted by me, for the BCSE308P – Computer Networks "
        "to VIT is a record of bonafide work carried out by me under the supervision of Dr. S. Ebenezer Juliet.")
P("I further declare that the work reported in this report has not been submitted and will not be submitted, either in part "
  "or in full, for any other courses in this institute or any other institute or university.", after=60)
P("Place : Vellore", align="left", after=2)
p = P("Date : ", align="left", after=2); placeholder(p, "[dd-mm-yyyy]")
p = P("", align="right", after=2); p.add_run("Signature of the Candidate(s)")
page_break()

# =============================================================================== TOC (field with static fallback)
TOC_ENTRIES = ["1. Introduction", "2. Problem Statement", "3. Objectives", "4. Literature Review and Existing System",
               "5. Abstract of the Proposed Project Work", "6. Workflow of the Proposed Project Work / System Design / System Architecture",
               "7. Contribution of Individual Team Member", "8. Modules Explanation", "9. Project Implementation",
               "    a. Experimental Setup with Hardware and Software Requirements", "    b. Test Cases", "    c. Results of Execution",
               "    d. Performance Analysis – Comparison of Methods / Protocols", "    e. Critical Evaluation / Optimization",
               "10. Results and Discussion", "    a. Findings and Inferences", "    b. Innovation, Novelty, and Evidence",
               "11. Conclusion", "12. Limitations and Future Enhancement", "References"]
p = P("TABLE OF CONTENTS", align="center", size=14, bold=True, after=14)
for i, e in enumerate(TOC_ENTRIES):
    para = doc.add_paragraph()
    para.paragraph_format.space_after = Pt(4)
    if i == 0:
        for typ in ("begin", "instr", "separate"):
            r = para.add_run()
            if typ == "instr":
                it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = ' TOC \\o "1-2" \\h \\z \\u '; r._r.append(it)
            else:
                f = OxmlElement("w:fldChar"); f.set(qn("w:fldCharType"), typ); r._r.append(f)
    para.add_run(e)
    if i == len(TOC_ENTRIES) - 1:
        r = para.add_run(); f = OxmlElement("w:fldChar"); f.set(qn("w:fldCharType"), "end"); r._r.append(f)
settings = doc.settings.element
uf = OxmlElement("w:updateFields"); uf.set(qn("w:val"), "true"); settings.append(uf)
page_break()

# footer page numbers
fp = sec.footer.paragraphs[0]
fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
field(fp, " PAGE ")

# =============================================================================== 1 Introduction
H1("1. Introduction")
P("Almost every networked application is built on the client–server model: a long-running server process listens on a well-known "
  "address, and clients initiate requests to it. The programming interface that exposes this model to applications is the socket API. "
  "Sockets give a program direct access to the two transport protocols of the Internet: the Transmission Control Protocol (TCP), "
  "which provides a reliable, ordered byte stream at the cost of connection set-up and per-connection state, and the User Datagram "
  "Protocol (UDP), which provides unreliable, unordered datagrams with almost no overhead. Choosing between them, and building "
  "whatever reliability the application needs on top, is a central design decision in network programming.")
P("This project builds a **lightweight client–server network service platform** that exposes three classic services from a single "
  "server process: **Echo** (TCP and UDP), **File Transfer** (TCP) and **DNS** (UDP, real RFC 1035 wire format). The platform "
  "serves many clients at once, writes a structured communication log, detects and survives communication failures, and is "
  "measured on response time, throughput, reliability, packet loss and resource utilisation. Each service uses the transport that "
  "suits it, so the same code base can be used to compare TCP and UDP behaviour directly.")
P("The implementation is written in Python using only the standard library (`socket`, `threading`, `struct`, `hashlib`, `logging`), "
  "so every protocol detail, from the 4-byte frame header of the echo service to the compression pointers in DNS responses, is "
  "visible in the source rather than hidden in a framework.")

# =============================================================================== 2 Problem
H1("2. Problem Statement")
P("*Design and implement a lightweight client–server network service platform using TCP and UDP socket programming. The system "
  "should provide Echo, File Transfer, and DNS services. The server must support multiple concurrent clients, maintain communication "
  "logs, detect communication failures, and provide appropriate error handling. Evaluate the system based on response time, "
  "throughput, reliability, packet loss, and resource utilization.*  (Problem III, Innovative Practices – Computer Networks Lab, "
  "Fall Semester 2026-2027.)")
P("The problem therefore has a construction half (four requirements on the server: three services, concurrency, logging, "
  "failure detection and error handling) and an evaluation half (five metrics). The report addresses both: Sections 6–8 describe the "
  "construction, and Sections 9–10 present the evidence for each metric.")

# =============================================================================== 3 Objectives
H1("3. Objectives")
bullets([
    "Implement an **Echo service** over TCP (length-prefixed frames) and over UDP (datagram echo).",
    "Implement a **File Transfer service** over TCP supporting list, upload, download and status, with end-to-end integrity checking.",
    "Implement a **DNS service** over UDP that speaks the real DNS wire format (A, AAAA, CNAME, MX, TXT) and interoperates with standard tools.",
    "Serve **multiple concurrent clients** without cross-talk, with configurable limits.",
    "Maintain **communication logs** (connections, requests, failures, durations, byte counts) and live server statistics.",
    "**Detect communication failures** (crashed or hung peers, idle clients, truncated transfers, malformed packets, lost datagrams) and handle each with an appropriate error response or recovery action.",
    "**Evaluate** the platform on response time, throughput, reliability, packet loss and resource utilisation, using repeatable experiments.",
    "Go beyond measurement: identify bottlenecks, implement optimisations and quantify their effect.",
])

# =============================================================================== 4 Literature
H1("4. Literature Review and Existing System")
P("**Transport protocols.** TCP (RFC 793, updated by RFC 9293) provides reliable, ordered delivery using acknowledgements, "
  "retransmission and congestion control [1]. UDP (RFC 768) adds only ports and a checksum to IP, leaving loss recovery to the "
  "application [2]; RFC 8085 gives guidelines for applications that must therefore implement their own timeouts and back-off [3]. "
  "Stevens, Fenner and Rudoff [4] and Kurose and Ross [5] describe the socket API and the design trade-offs between iterative, "
  "forking, threaded and event-driven servers.")
P("**Existing services.** The Echo protocol (RFC 862) [6] is the traditional diagnostic service. File transfer is standardised by "
  "FTP (RFC 959) [7], which uses separate control and data connections, and by TFTP (RFC 1350) [8], which runs over UDP with "
  "lock-step acknowledgements and is therefore simple but slow. DNS is specified in RFC 1034 and RFC 1035 [9], [10]; production "
  "implementations include BIND and dnsmasq. Scalability of connection-per-thread servers is discussed in Kegel’s “C10K problem” [11].")
table(["System", "Transport", "Strength", "Gap relative to this project"],
      [["inetd / xinetd echo", "TCP+UDP", "Standard, tiny", "No logging of failures, no metrics, no integrity checks, single service"],
       ["vsftpd / FTP (RFC 959)", "TCP (2 connections)", "Feature rich", "Heavy: separate data channel, authentication, active/passive modes; not meant for protocol study"],
       ["TFTP (RFC 1350)", "UDP", "Simple", "Lock-step ACKs limit throughput; no integrity beyond UDP checksum"],
       ["BIND / dnsmasq", "UDP+TCP", "Complete DNS", "Large code bases; not instrumented for loss/latency experiments"],
       ["Python `socketserver` examples", "TCP/UDP", "Short", "Single service, no framing, no limits, no failure handling or evaluation"],
       ["**This platform**", "TCP+UDP", "Four services, one process, shared logging/stats/failure handling, loss injection, reproducible benchmarks", "Educational scale: no TLS, authentication or recursion (Section 12)"]],
      widths=[1.45, 1.0, 1.7, 2.1], cap="Comparison with existing systems", size=9)
P("**Gap addressed.** Reference implementations are either too large to read in an afternoon or too small to evaluate. The "
  "platform sits between them: small enough to inspect completely, complete enough (concurrency, logging, failure handling, loss "
  "injection, instrumentation) to support quantitative experiments on the five required metrics.")

# =============================================================================== 5 Abstract
H1("5. Abstract of the Proposed Project Work")
P(f"The project delivers a client–server platform in which one server process offers four services: a TCP echo service with "
  f"length-prefixed framing, a UDP echo service, a TCP file-transfer service with SHA-256 integrity verification in both directions "
  f"and atomic stores, and an authoritative UDP DNS service that implements RFC 1035 message encoding and decoding and interoperates with "
  f"the standard `dig` tool. TCP services use one thread per connection with idle timeouts and a client cap; UDP services use a "
  f"receive loop and a worker pool. All services write a structured log and update shared counters that clients can query "
  f"at run time. The server detects timeouts, resets, truncated messages, oversize frames, malformed DNS packets, bad file names and "
  f"checksum mismatches, and recovers from each without affecting other clients. The client library adds bounded connect retries with "
  f"exponential back-off and UDP retransmission with per-attempt timeouts.")
P(f"Evaluation on an {M['cpu']} machine ({M['cores']} cores, {M['ram_gb']:.0f} GB) over the loopback interface shows: TCP echo "
  f"round-trip times of {f3(tcp1k['p50_ms'])} ms (1 KiB, median); a sustained {f0(sum(plateau) / len(plateau))} echo requests per second "
  f"(peak {f0(peak['req_per_s'])}) with 0 errors for 1 to 400 concurrent clients; file transfers at about {f1(up100 / 1000)} GB/s; "
  f"{f0(max(r['req_per_s'] for r in E4Q))} DNS queries per second; UDP loss measurements that match the analytical model 1−(1−p)² to within "
  f"{MAXDEV_RAW:.1f} percentage points, and a retransmission policy that raises DNS success from {pc(dns_at(loss20, 0)['success'])} to "
  f"{pc(dns_at(loss20, 2)['success'])} at 20% per-datagram loss; and {f1(idle2000['kb_per_conn'])} KB of memory per open connection. "
  f"A digest cache raised single-client download throughput from {f1(dl_base100 / 1000)} to {f1(dl_dc100 / 1000)} GB/s, and combined with "
  f"`sendfile()` it cut server CPU at 8 concurrent downloads from {f0(c8[0]['cpu_pct'])}% to {f0(c8[3]['cpu_pct'])}%. An adaptive, RFC 6298-style retransmission timer cut "
  f"the response time of lossy DNS queries by {min(RTO_SPEEDUP):.1f}–{max(RTO_SPEEDUP):.1f}× at the same success rate. The suite of "
  f"{T['total']} automated tests ({tests_pass} passing) covers functional, fault-injection and security cases.")

# =============================================================================== 6 Workflow / architecture
H1("6. Workflow of the Proposed Project Work / System Design / System Architecture")
H2("6.1 Architecture")
figure("fig_architecture.png", 6.1, "System architecture: four services in one server process, shared logging, statistics and failure handling")
table(["Service", "Transport / port", "Request → response", "Concurrency model", "Reliability mechanism"],
      [["Echo", "TCP 9001", "`[len:4][payload]` → same frame", "1 thread per connection", "TCP + frame length check; idle timeout"],
       ["File", "TCP 9002", "text command line (+ binary body) → `OK`/`ERR` line (+ body)", "1 thread per connection", "SHA-256 end to end; atomic rename; abort cleanup"],
       ["UDP echo", "UDP 9003", "datagram → same datagram", "receive loop + 16 workers", "client timeout and retransmission"],
       ["DNS", "UDP 9053", "RFC 1035 query → response (rcode, answers)", "receive loop + 16 workers", "client timeout, retransmission, same transaction ID"]],
      widths=[0.7, 0.9, 1.9, 1.3, 1.6], cap="Service summary", size=9)
H2("6.2 Request workflows")
P("**TCP echo.** (1) Client connects (bounded retries with exponential back-off). (2) Client sends a 4-byte big-endian length followed by the "
  "payload. (3) The server thread reads exactly that many bytes, rejects lengths above 16 MiB, and writes the frame back unchanged. "
  "(4) The client verifies the reply equals the request and records the round-trip time. (5) Either side may close; the server logs "
  "the session summary (requests, bytes, duration, outcome).")
P("**File upload (PUT).** (1) Server validates the name against `^[A-Za-z0-9_][A-Za-z0-9._-]{0,127}$` and the size against the limit, "
  "then replies `READY`. (2) Client streams the bytes. (3) Server writes them to a hidden temporary file while hashing. (4) If the digest "
  "matches the announced one the file is moved into place with an atomic rename and the server replies `OK <size>`; otherwise it replies `ERR 422` "
  "and discards the data. (5) If the client disappears mid-transfer, the temporary file is removed and an `aborted` event is logged. "
  "**Download (GET)** is symmetric: the server announces `OK size sha256`, streams the body, and the client verifies the digest.")
P("**DNS query.** (1) Client builds a query with a fresh transaction ID and sends it. (2) If no matching reply arrives within the timeout it "
  "retransmits the *same* packet, up to the retry budget; replies with a different ID are ignored. (3) The server parses the header and question "
  "(rejecting compression loops, truncated names and QDCOUNT≠1 with FORMERR), looks the name up case-insensitively, follows CNAME chains up to "
  "8 hops, and answers NOERROR (with or without data) or NXDOMAIN; unknown names can optionally be forwarded to an upstream resolver and cached.")
H2("6.3 Failure-handling design")
table(["Failure", "Detected by", "Server action", "Client action"],
      [["Client crashes / disconnects mid-message", "`recv` returns 0 with a partial message", "Log `aborted`, delete temp file, free slot", "—"],
       ["Client silent (idle)", "Socket idle timeout (default 30 s)", "Log `idle_timeout`, close", "Next call raises `ConnectionClosed`"],
       ["Connection reset / broken pipe", "`ConnectionResetError`, `BrokenPipeError`", "Log `reset`, clean up", "Retry or report"],
       ["Server crashes", "Peer EOF / RST", "—", "`ConnectionClosed` within microseconds"],
       ["Server hangs", "No reply within socket timeout", "—", "`TimeoutError` after the configured timeout"],
       ["Server not running", "`connect` refused", "—", "3 attempts with 0.2 s, 0.4 s back-off, then `ConnectionError`"],
       ["Too many clients", "Semaphore of `--max-clients`", "`ERR 503 server busy`, log `rejected`", "Error surfaced to caller"],
       ["Oversize frame / line / file", "Length field checked before allocation", "Close (echo), `ERR 400/413` (file)", "Error surfaced"],
       ["Path traversal, bad names", "Whitelist regular expression", "`ERR 400`, session continues", "`ServiceError`"],
       ["Corrupted upload/download", "SHA-256 mismatch", "`ERR 422`, data discarded", "`ProtocolError`"],
       ["Malformed DNS packet", "Parser raises `DNSFormatError`", "FORMERR reply (or drop if no ID), log `malformed`", "rcode shown"],
       ["Lost UDP datagram", "Client timeout", "—", "Retransmit up to the retry budget, then `socket.timeout`"]],
      widths=[1.55, 1.6, 1.7, 1.55], cap="Failure scenarios and handling", size=9)
H2("6.4 Concurrency and logging design")
P("For each TCP service an *accept thread* only calls `accept()` and enqueues the socket, so the kernel listen queue drains as fast as "
  "possible; a *dispatcher thread* then acquires a slot from a bounded semaphore (`--max-clients`) and starts one handler thread per connection. "
  "For UDP, a *receive thread* reads datagrams and hands them to a `ThreadPoolExecutor`. A single `Stats` object (lock-protected counters) is "
  "shared by all services and can be read remotely with the `STATS` command. Logging uses one `key=value` line per event "
  "(`svc=file client=127.0.0.1:55635 event=aborted detail=...`), written to a rotating file and to stderr, which makes the logs easy to grep and parse.")

# =============================================================================== 7 Contribution
H1("7. Contribution of Individual Team Member")
table(["Name", "Register No.", "Contribution"],
      [[NAME, REG, "Sole author. Requirements analysis; protocol design (framing, file protocol, DNS wire format); implementation of server, client library and CLI; "
                   f"test suite ({T['total']} tests); benchmark harness; execution of all experiments; analysis; optimisation work; this report."]],
      widths=[1.4, 1.3, 3.6], cap="Individual contribution (edit if the project was done by a team)", size=10)

# =============================================================================== 8 Modules
H1("8. Modules Explanation")
table(["Module (file)", "Responsibility", "Key design points"],
      [["`protocol.py`", "Framing helpers shared by TCP services", "`recv_exact` loops until n bytes arrive and distinguishes a clean close from a mid-message close (`ConnectionClosed.partial`); 16 MiB frame cap; `sha256_file`"],
       ["`dns_wire.py`", "RFC 1035 encoder/decoder", "Label encoding, compression-pointer decoding with loop limit, A/AAAA/CNAME/MX/TXT RDATA, header flags (AA, TC, RD, RA, rcode), 512-byte UDP limit with TC bit"],
       ["`server.py` – `Stats`, `Session`", "Counters, per-connection accounting", "Thread-safe counters, active/peak connections, `getrusage` CPU and peak RSS"],
       ["`server.py` – `TCPService`", "Accept/dispatch/serve lifecycle for TCP services", "Accept queue decoupled from thread start, semaphore-bounded clients, idle timeout, TCP keep-alive, uniform exception → log mapping"],
       ["`server.py` – `EchoService`", "TCP echo", "Frame in, same frame out"],
       ["`server.py` – `FileService`", "File transfer", "LIST/GET/PUT/STAT/STATS/PING/QUIT; name whitelist; size limit; temp-file + atomic rename; SHA-256 digest cache; chunked or `sendfile` sending"],
       ["`server.py` – `UDPService`, `UDPEchoService`", "UDP receive loop and worker pool", "Per-datagram simulated loss (`--loss`) applied on receive and on send"],
       ["`server.py` – `Zone`, `DNSService`", "Authoritative DNS", "JSON zone, case-insensitive lookup, CNAME chasing, NXDOMAIN/NODATA, FORMERR/NOTIMP/REFUSED, optional forwarding with TTL cache"],
       ["`client.py`", "Client library and CLI", "`tcp_connect` (retry + back-off), `EchoClient`, `FileClient` (verified GET/PUT), `udp_request` (retransmit, ID matching), `RTTEstimator` (adaptive RTO), `dns_query`"],
       ["`tests/`", f"{T['total']} automated tests", "Functional, concurrency, fault injection, security; `run_tests.py` records results as JSON"],
       ["`bench/run_benchmarks.py`", "Evaluation harness", "Launches the server as a separate process, multi-process load generators, RSS sampling, CPU from `getrusage` deltas"]],
      widths=[1.55, 1.5, 3.35], cap="Module overview", size=9)
H2("8.1 Protocol specifications")
table(["Command / frame", "Format", "Success reply", "Errors"],
      [["Echo frame", "`len:uint32 BE` + payload", "same frame", "len > 16 MiB → connection closed"],
       ["`PING`", "`PING\\n`", "`OK pong`", "—"],
       ["`LIST`", "`LIST\\n`", "`OK <n>\\n` + n bytes JSON", "—"],
       ["`STAT name`", "`STAT report.bin\\n`", "`OK <size> <sha256>`", "`400` bad name, `404`"],
       ["`GET name`", "`GET report.bin\\n`", "`OK <size> <sha256>\\n` + size bytes", "`400`, `404`"],
       ["`PUT name size sha`", "`PUT a.bin 1000 <64 hex>\\n` + bytes after `READY`", "`READY`, then `OK <size>`", "`400`, `413` too large, `422` checksum"],
       ["`STATS`", "`STATS\\n`", "`OK <n>\\n` + JSON counters", "—"],
       ["`QUIT`", "`QUIT\\n`", "`OK bye`", "—"],
       ["Any unknown", "—", "—", "`ERR 400 unknown or malformed command`; line > 1 KiB closes"],
       ["Busy", "—", "—", "`ERR 503 server busy` then close"]],
      widths=[1.2, 2.3, 1.6, 1.3], cap="TCP application protocols", size=9)

# =============================================================================== 9 Implementation
H1("9. Project Implementation")
H2("a. Experimental Setup with Hardware and Software Requirements")
table(["Item", "Specification"],
      [["Machine", f"{M['cpu']}, {M['cores']} cores, {M['ram_gb']:.0f} GB RAM"],
       ["Operating system", M["platform"]],
       ["Language / runtime", f"Python {M['python']} (standard library only for the platform, tests and benchmarks)"],
       ["Report tooling", "matplotlib (figures), python-docx (this document)"],
       ["Network", "Loopback interface (127.0.0.1): server, clients and load generators are separate processes on one host"],
       ["External tool for interoperability check", "`dig` (BIND utilities) as an independent DNS client"],
       ["Kernel settings observed", f"`kern.ipc.somaxconn` = {E6B['somaxconn']}; file-descriptor soft limit raised by the server to 8192"],
       ["Minimum requirements to run", "Any OS with Python 3.9+; ≈30 MB RAM + ≈54 KB per open connection; no third-party packages"]],
      widths=[2.0, 4.4], cap="Experimental setup", size=10)
P("**Methodology.** Each experiment starts a fresh server process (`bench/run_benchmarks.py`) so that results are independent. "
  "Latencies are measured on the client with `time.perf_counter()` around one request/response. Load generators run in four "
  "separate processes to avoid a single-interpreter bottleneck on the client side, with clients started 2 ms apart to model realistic arrivals "
  "(the reason is explained in Section 9.e). Server CPU is the difference of `getrusage` user+system time divided by wall time, so 100% means one "
  "fully busy core; server memory is sampled from `ps` every 100 ms. Single-client results are means of 3 repetitions; concurrent "
  "file-transfer results are single runs, so differences below roughly 15% should be treated as noise. Packet loss is *simulated* inside the "
  "server (each datagram received or sent is dropped with probability p), because loopback never drops packets; this makes the "
  "loss experiments exactly repeatable and checkable against theory.")

H2("b. Test Cases")
P(f"The automated suite contains {T['total']} test cases; **{tests_pass} passed, {T['failures']} failed, {T['errors']} errors, "
  f"{T['skipped']} skipped** in {T['seconds']} s (`python3 tests/run_tests.py`).")
trows = []
for i, t in enumerate(T["tests"], 1):
    cls, fn = t["id"].split(".")[-2:]
    trows.append([f"TC{i:02d}", f"{cls[4:]}: `{fn[5:]}`", t["description"], t["status"]])
table(["ID", "Area / test", "Description and expected behaviour", "Result"], trows, widths=[0.5, 1.7, 3.55, 0.65],
      cap="Test cases and results", size=8.5, align_num=False)

H2("c. Results of Execution")
P("The screenshots below were captured from an actual run of `demo.sh` (server in one terminal, clients in another).")
figure("term_server.png", 5.6, "Starting the server: four services listening")
figure("term_echo.png", 5.2, "TCP and UDP echo with measured round-trip time")
figure("term_file.png", 5.6, "File transfer: upload, list, download with SHA-256 verification, byte-for-byte comparison, and rejected requests")
figure("term_dns.png", 5.8, "DNS service queried with the project client and with the independent `dig` tool (CNAME chain, MX, NXDOMAIN)")
figure("term_concurrency.png", 4.4, "Twenty parallel clients and the live server statistics returned by the `STATS` command")
figure("term_failure.png", 5.4, "Fault injection: aborted upload, path traversal and malformed DNS packet (FORMERR, rcode 1)")
figure("term_log.png", 6.2, "Server log excerpt: uploads, downloads, DNS queries with microsecond service times, a WARNING for the aborted upload and for the malformed packet")
figure("term_tests.png", 6.3, f"Automated test-suite output ({tests_pass}/{T['total']} passing)")

H2("d. Performance Analysis – Comparison of Methods / Protocols")
H3("d.1 Response time versus message size (TCP versus UDP)")
rows = []
for s_ in E1T:
    u = [x for x in E1U if x["size"] == s_["size"]]
    rows.append([f"{s_['size']:,}", f3(s_["p50_ms"]), f3(s_["p95_ms"]), f3(s_["p99_ms"]), f0(s_["goodput_MBps"]),
                 f3(u[0]["p50_ms"]) if u else "—", f3(u[0]["p95_ms"]) if u else "—"])
table(["Payload (B)", "TCP p50 (ms)", "TCP p95 (ms)", "TCP p99 (ms)", "TCP goodput (MB/s)", "UDP p50 (ms)", "UDP p95 (ms)"], rows,
      widths=[0.9, 0.85, 0.85, 0.85, 1.05, 0.85, 0.85], cap="Echo round-trip time on a persistent connection (1,000 requests per size; 100 for ≥256 KiB). UDP sizes above 8 KiB exceed the default loopback datagram limit.", size=9, label="echo_lat")
P(f"One-shot cost: opening a new TCP connection for every request costs {f3(E1N['mean_ms'])} ms on average (p95 {f3(E1N['p95_ms'])} ms) versus "
  f"{f3(udp16['p50_ms'])} ms for a 16-byte UDP request, i.e. UDP is about {E1N['mean_ms'] / udp16['p50_ms']:.1f}× faster for an isolated exchange because there is no three-way handshake or teardown.")
figure("fig_latency.png", 6.4, "Response time (left) and goodput (right) versus payload size")

H3("d.2 Concurrency, throughput and server utilisation")
table(["Clients", "Requests/s", "p50 (ms)", "p95 (ms)", "p99 (ms)", "Errors", "Server CPU (%)", "Server RSS (MB)"],
      [[r["clients"], f0(r["req_per_s"]), f2(r["p50_ms"]), f2(r["p95_ms"]), f2(r["p99_ms"]), r["errors"], f0(r["cpu_pct"]), f1(r["rss_peak_mb"])] for r in E2],
      widths=[0.65, 0.9, 0.7, 0.7, 0.7, 0.6, 1.0, 1.0], cap="TCP echo (1 KiB messages) under increasing concurrency", size=9, label="conc")
figure("fig_concurrency.png", 6.5, "Throughput, response-time percentiles and server CPU versus number of concurrent clients")

H3("d.3 File transfer throughput")
rows = [[f"{b['size_mb']} MB", f0(b["upload_MBps"]), f0(b["download_MBps"]), f0(s_["download_MBps"]), f0(d["download_MBps"]), f0(e["download_MBps"])]
        for b, s_, d, e in zip(E3, SF, DC, DS)]
table(["File size", "Upload (MB/s)", "Download baseline", "Download sendfile", "Download digest cache", "Download cache + sendfile"], rows,
      widths=[0.8, 0.95, 1.05, 1.05, 1.2, 1.35], cap="Single-client file transfer throughput (mean of 3 runs; SHA-256 verified on every transfer)", size=9, label="file_single")
rows = [[c["clients"]] + sum(([f0(v[i]["aggregate_MBps"]), f0(v[i]["cpu_pct"])] for v in (E3B, SFC, DCC, DSC)), []) for i, c in enumerate(E3B)]
table(["Concurrent downloads", "Baseline MB/s", "CPU %", "sendfile MB/s", "CPU %", "Digest cache MB/s", "CPU %", "Cache+sendfile MB/s", "CPU %"], rows,
      widths=[0.85, 0.7, 0.5, 0.7, 0.5, 0.8, 0.5, 0.85, 0.5], cap="Aggregate throughput and server CPU for N simultaneous 20 MB downloads", size=8.5, label="file_conc")
figure("fig_files.png", 6.6, "File-transfer throughput and server CPU for the four server variants")

H3("d.4 DNS service")
rows = [[k, f3(v["mean_ms"]), f3(v["p50_ms"]), f3(v["p95_ms"]), f3(v["p99_ms"])] for k, v in E4L.items()]
table(["Query type", "Mean (ms)", "p50 (ms)", "p95 (ms)", "p99 (ms)"], rows, widths=[1.8, 1, 1, 1, 1],
      cap="DNS response time, 3,000 sequential queries per type (client opens a new socket per query)", size=9.5, label="dns_lat")
table(["Concurrent resolvers", "Queries/s", "p95 (ms)", "Errors", "Server CPU (%)"],
      [[r["clients"], f0(r["req_per_s"]), f2(r["p95_ms"]), r["errors"], f0(r["cpu_pct"])] for r in E4Q],
      widths=[1.5, 1.1, 1, 0.8, 1.2], cap="DNS query rate under concurrency", size=9.5)
figure("fig_dns.png", 6.3, "DNS response time by query type, and query rate / p95 latency versus concurrency")

H3("d.5 Packet loss and reliability")
rows = []
for r in E5:
    rows.append([pc(r["loss_per_datagram"]), pc(r["udp_theory_loss"]), pc(r["udp_observed_loss"])]
                + [f"{pc(dns_at(r, k)['success'])} ({pc(dns_at(r, k)['theory_success'])})" for k in (0, 1, 2, 4)]
                + [pc(r["tcp_echo_success"])])
table(["Loss per datagram p", "Round-trip loss, theory", "Round-trip loss, measured", "DNS success, 0 retries (theory)", "1 retry", "2 retries", "4 retries", "TCP echo success"],
      rows, widths=[0.7, 0.75, 0.8, 1.05, 0.9, 0.9, 0.9, 0.7],
      cap=f"UDP loss and DNS success rate versus retransmission budget ({NQ:,} queries per cell; 2,000 raw datagrams per row; 50 ms timeout)", size=8.5, label="loss_tab")
rows = [[pc(r["loss_per_datagram"])] + [f"{f1(dns_at(r, k)['mean_ms_success'])} / {f1(dns_at(r, k)['p95_ms_success'])}" for k in (0, 1, 2, 4)] +
        [f"{dns_at(r, 2)['mean_attempts']:.2f}"] for r in E5]
table(["Loss p", "0 retries mean / p95 (ms)", "1 retry", "2 retries", "4 retries", "Mean attempts (2 retries)"], rows,
      widths=[0.7, 1.35, 1.05, 1.05, 1.05, 1.2], cap="Response time of successful DNS queries (mean / p95, ms)", size=9)
figure("fig_loss.png", 6.7, "Measured versus theoretical UDP loss (left); DNS success versus retry budget (centre); latency cost of retransmission (right)")

H3("d.6 Resource utilisation and scalability limits")
table(["Idle connections", "Server RSS (MB)", "Server threads", "KB per connection", "Echo p50 (ms)", "Echo p95 (ms)"],
      [[r["idle_connections"], f1(r["rss_mb"]), r["threads"], f1(r["kb_per_conn"]) if r["idle_connections"] else "—", f3(r["echo_p50_ms"]), f3(r["echo_p95_ms"])] for r in E6],
      widths=[1.0, 1.0, 1.0, 1.1, 1.0, 1.0], cap="Memory and thread cost of idle connections, and latency of one active client alongside them", size=9.5, label="idle")
figure("fig_resources.png", 6.7, "Server memory under load (left), memory and threads versus idle connections (centre), and latency of an active client (right)")
bt = [t_["connections_before_failure"] for t_ in E6B["trials"]]
P(f"**Connection bursts.** Opening connections back-to-back with no pause, the kernel accept queue overflowed after {min(bt)}–{max(bt)} connections "
  f"in three trials (`ConnectionResetError`), consistent with `kern.ipc.somaxconn` = {E6B['somaxconn']}; the same server accepted 2,000 connections "
  f"when they arrived at a realistic pace (10 connections per 3 ms).")

H3("d.7 Failure-detection time")
table(["Scenario", "How the failure appears", "Detection time"],
      [["Server process hung (SIGSTOP)", f"`{E7['hung_server']['exception']}` at client socket timeout of {E7['hung_server']['client_timeout_s']:.1f} s", f"{f3(E7['hung_server']['detected_after_s'])} s"],
       ["Server process killed (SIGKILL)", f"`{E7['crashed_server']['exception']}` (kernel closes the socket)", f"{E7['crashed_server']['detected_after_s'] * 1000:.2f} ms"],
       ["Server not listening (TCP connect)", "`ConnectionError` after 3 attempts, back-off 0.2 s + 0.4 s", f"{f3(E7['connect_refused']['detected_after_s'])} s"],
       ["Server not listening (UDP)", "`socket.timeout` after 3 attempts × 0.2 s – UDP gives no error", f"{f3(E7['udp_dead_server']['detected_after_s'])} s"],
       ["Idle client (server side)", f"Server closes after configured idle timeout of {E7['server_idle_drop']['configured_idle_timeout_s']:.1f} s", f"{f3(E7['server_idle_drop']['closed_after_s'])} s"]],
      widths=[2.0, 3.2, 1.2], cap="Time to detect communication failures", size=9.5, label="fail_det")

H3("d.8 TCP versus UDP: summary comparison")
table(["Criterion", "TCP services (echo, file)", "UDP services (echo, DNS)", "Evidence"],
      [["Connection set-up", "3-way handshake per connection", "No handshake", f"One-shot request: {f3(E1N['mean_ms'])} ms (TCP, new connection) vs {f3(udp16['p50_ms'])} ms (UDP)"],
       ["Steady-state latency (1 KiB)", f"{f3(tcp1k['p50_ms'])} ms p50", f"{f3(udp1k['p50_ms'])} ms p50", "UDP client creates a socket per request and the server hands off to a worker thread (Section 10.a)"],
       ["Large transfers", "Stream, flow/congestion control: 1.0–1.7 GB/s", "Datagram-bounded; not suitable for bulk data", REF["file_single"]],
       ["Loss", "Hidden by kernel retransmission (delay, not failure)", f"Visible: {pc(loss20['udp_observed_loss'])} round-trip loss at p = 20%", REF["loss_tab"]],
       ["Reliability burden", "Transport", f"Application: {pc(dns_at(loss20, 0)['success'])} → {pc(dns_at(loss20, 2)['success'])} success with 2 retries", REF["loss_tab"]],
       ["Server state / memory", f"Thread + socket: ≈{f0(idle2000['kb_per_conn'])} KB per connection", "No per-client state (worker pool only)", REF["idle"]],
       ["Failure detection", "Immediate on close/reset; timeout if hung", "Only timeouts", REF["fail_det"]]],
      widths=[1.2, 1.7, 1.7, 1.8], cap="TCP versus UDP in this platform", size=9)

H2("e. Critical Evaluation / Optimization")
P("This section records what the measurements revealed as weaknesses, what was changed in response, and what remains.")
P(f"**Finding and mitigation 1 – connection bursts and the accept path.** The first full benchmark lost one of 400 clients (50 requests) and a test that opened 500 idle "
  f"connections back-to-back was reset by the kernel. The listen queue holds at most `kern.ipc.somaxconn` = {E6B['somaxconn']} pending connections, so a client that connects faster than the server "
  f"can accept overflows it. Two server-side changes were tried: moving thread creation out of the accept loop into a separate dispatcher thread (kept, because it keeps `accept()` responsive at "
  f"no cost), and shortening the interpreter’s GIL switch interval (no measurable effect, removed). Neither removed the limit: back-to-back connects still fail after {min(bt)}–{max(bt)} "
  f"connections on this machine, and the exact figure varies from run to run. The limit is therefore documented as a finding, the benchmark ramps clients in at a realistic pace, and the proper fixes are "
  f"listed in Section 12 (larger `somaxconn`/backlog, pre-spawned worker pool or event loop). With the ramp, 400 concurrent clients completed with 0 errors and 2,000 idle connections were held.")
P(f"**Optimisation 1 – SHA-256 digest cache.** Uploads ran at {f0(up100)} MB/s but downloads at only {f0(dl_base100)} MB/s (100 MB file). The reason is structural: to announce the digest "
  f"before the body, the server read the whole file once to hash it and a second time to send it. Caching the digest keyed by (path, mtime, size) – "
  f"populated when a verified upload completes and invalidated automatically when the file changes (test TC{[i for i, t in enumerate(T['tests'], 1) if 'digest_cache' in t['id']][0]:02d}) – "
  f"raised single-client downloads to {f0(dl_dc100)} MB/s (**+{(dl_dc100 / dl_base100 - 1) * 100:.0f}%**), equal to the upload speed.")
P(f"**Optimisation 2 – `sendfile()`.** Replacing the `read()`+`send()` loop with `socket.sendfile()` gave no gain for a single client "
  f"({f0(dl_sf100)} versus {f0(dl_base100)} MB/s), which is expected when the hash pass dominates. Its benefit shows under concurrency, where the saved user-space copies "
  f"reduce CPU: with 8 concurrent downloads the baseline needed {f0(c8[0]['cpu_pct'])}% CPU for {f0(c8[0]['aggregate_MBps'])} MB/s, whereas cache + `sendfile()` "
  f"needed {f0(c8[3]['cpu_pct'])}% CPU for {f0(c8[3]['aggregate_MBps'])} MB/s – roughly {c8[3]['aggregate_MBps'] / c8[0]['aggregate_MBps']:.1f}× the throughput at "
  f"{c8[3]['cpu_pct'] / c8[0]['cpu_pct'] * 100:.0f}% of the CPU. (Single run; the direction was the same in every repeated run, and with 16 clients the gain is "
  f"{c16[3]['aggregate_MBps'] / c16[0]['aggregate_MBps']:.1f}×.)")
P(f"**Weakness 1 – single-core ceiling.** Echo throughput plateaus at {f0(min(plateau))}–{f0(max(plateau))} requests/s from 10 to 100 clients while server CPU is "
  f"≈{f0(by(E2, 'clients', 50)['cpu_pct'])}% – one saturated core. CPython’s Global Interpreter Lock lets only one thread execute Python bytecode at a time, so adding "
  f"clients adds queueing delay, not capacity: p95 latency grows almost linearly (Little’s law: {c100['clients']} clients / {f0(c100['req_per_s'])} req/s = "
  f"{c100['clients'] / c100['req_per_s'] * 1000:.2f} ms predicted mean versus {f2(c100['mean_ms'])} ms measured). At 400 clients throughput falls to {f0(c400['req_per_s'])} req/s "
  f"({(1 - c400['req_per_s'] / (sum(plateau) / len(plateau))) * 100:.0f}% below the plateau) because 400 runnable threads add scheduling overhead. The fix would be several worker processes sharing the listening port "
  f"(`SO_REUSEPORT`) or an event-loop design; it was not implemented because it would change the thread-per-connection teaching model.")
P(f"**Weakness 2 – memory per connection.** Thread-per-connection costs ≈{f1(idle2000['kb_per_conn'])} KB per client "
  f"({f1(idle2000['rss_mb'])} MB at 2,000 connections) and `kern.num_taskthreads` is 4,096 on this machine, so the design tops out near 4,000 simultaneous clients. Latency of an active client "
  f"is hardly affected by idle ones (p50 {f3(idle0['echo_p50_ms'])} → {f3(idle2000['echo_p50_ms'])} ms), so this is a capacity limit, not a speed limit.")
P(f"**Optimisation 3 – adaptive retransmission timeout.** The loss experiments used a fixed 50 ms per-attempt timeout, so every lost datagram cost a full 50 ms "
  f"even though the round trip takes about 0.1 ms; at 40% loss with 4 retries the mean response was {f1(RTO40['fixed']['mean_ms'])} ms. The client now has an optional "
  f"`RTTEstimator` that follows RFC 6298 [12]: SRTT and RTTVAR are smoothed from RTT samples (α = 1/8, β = 1/4), RTO = SRTT + 4·RTTVAR clamped to [5 ms, 2 s], only replies to "
  f"first transmissions are sampled (Karn’s rule), and each retransmission doubles the timeout. Both variants were run against the same seeded loss sequence "
  f"({E9['queries']} DNS queries per cell, 4 retries).")
table(["Loss p", "Fixed: success", "Fixed: mean / p95 (ms)", "Adaptive: success", "Adaptive: mean / p95 (ms)", "Speed-up (mean)", "Learned RTO (ms)"],
      [[pc(r["loss_per_datagram"]), pc(r["fixed"]["success"]), f"{f2(r['fixed']['mean_ms'])} / {f2(r['fixed']['p95_ms'])}", pc(r["adaptive"]["success"]),
        f"{f2(r['adaptive']['mean_ms'])} / {f2(r['adaptive']['p95_ms'])}",
        f"{r['fixed']['mean_ms'] / r['adaptive']['mean_ms']:.1f}×" if r["loss_per_datagram"] > 0 else "—", f2(r["adaptive"]["median_rto_ms"])] for r in E9["rows"]],
      widths=[0.6, 0.8, 1.15, 0.85, 1.2, 0.8, 0.85], cap="Fixed versus adaptive retransmission timeout (DNS, 4 retries)", size=9, label="rto")
figure("fig_rto.png", 6.2, "Fixed 50 ms timeout versus adaptive RTO: response time (left) and success rate (right)")
P(f"The success rate is identical in every row, which confirms that reliability is set by the *number* of attempts, not by their timing. The response time, however, "
  f"falls by {min(RTO_SPEEDUP):.1f}–{max(RTO_SPEEDUP):.1f}×. On loopback the learned RTO sits at the 5 ms floor, so the floor (not the RTT) now dominates; on a real network with "
  f"tens of milliseconds of RTT the estimator would track the path instead. The trade-off is that a too-small floor causes spurious retransmissions on a path whose delay suddenly rises, "
  f"which is why RFC 6298 recommends a conservative minimum on the open Internet.")
P("**Threats to validity.** All traffic is on loopback, so there is no propagation delay, queueing in routers, MTU fragmentation or real loss; absolute numbers will be lower on a real network and "
  "the *relative* conclusions (TCP versus UDP behaviour, effect of retries, effect of optimisations) are the findings to rely on. Client and server share CPU cores, and the Python load generators "
  "contribute their own latency. Results come from one machine and one session.")

# =============================================================================== 10 Results
H1("10. Results and Discussion")
H2("a. Findings and Inferences")
bullets([
    f"**Response time.** A 1 KiB TCP echo takes {f3(tcp1k['p50_ms'])} ms (p50) and {f3(tcp1k['p99_ms'])} ms (p99) on a persistent connection; latency is flat up to 16 KiB and then grows with size as copying dominates "
    f"(1 MiB: {f3(tcp1m['p50_ms'])} ms, goodput {f0(tcp1m['goodput_MBps'])} MB/s). *Inference:* for small messages the cost is per-call overhead in the Python server, not bandwidth.",
    f"**TCP versus UDP.** On loopback a UDP echo (1 KiB: {f3(udp1k['p50_ms'])} ms) is slower than a TCP echo on an open connection ({f3(tcp1k['p50_ms'])} ms): the UDP client creates a socket for each request and the server "
    f"passes each datagram to a worker thread, whereas the TCP handler thread is already waiting in `recv`. This is an implementation effect; the protocol advantage of UDP appears only when connection set-up is included "
    f"({E1N['mean_ms'] / udp16['p50_ms']:.1f}× faster for a one-shot exchange). *Inference:* use UDP for short, independent request/response exchanges such as DNS, and TCP for sessions and bulk data.",
    f"**Concurrency.** Up to 400 simultaneous clients were served with 0 errors and no cross-talk. Throughput peaks at {f0(peak['req_per_s'])} req/s ({peak['clients']} clients) and plateaus around {f0(sum(plateau) / len(plateau))} req/s "
    f"with the server bound to one core; latency then rises linearly with the number of clients, as predicted by Little’s law.",
    f"**File transfer.** Uploads and downloads of 1–100 MB were all SHA-256 verified. Throughput is {f1(up100 / 1000)} GB/s for uploads and {f1(dl_base100 / 1000)} GB/s for downloads before optimisation, rising to "
    f"{f1(dl_dc100 / 1000)} GB/s afterwards. Values are far above any physical link because loopback is a memory copy; they measure the software stack, not a network.",
    f"**DNS.** Median service time is {f3(E4L['A (2 records)']['p50_ms'])} ms for an A query and {f3(E4L['CNAME chain']['p50_ms'])} ms for a two-level CNAME chain; the server itself spends only 15–30 µs per query "
    f"(from the log), so the rest is socket and scheduling cost. It sustains {f0(max(r['req_per_s'] for r in E4Q))} queries/s with 0 errors, and its answers are accepted by `dig`.",
    f"**Packet loss.** The measured round-trip loss follows 1−(1−p)² within {MAXDEV_RAW:.1f} percentage points at every setting. With no retransmission, 20% datagram loss means "
    f"{pc(1 - dns_at(loss20, 0)['success'])} of DNS queries fail; two retries restore {pc(dns_at(loss20, 2)['success'])} success (theory {pc(dns_at(loss20, 2)['theory_success'])}). "
    f"At 40% loss four retries still reach {pc(dns_at(loss40, 4)['success'])}. Across all 24 loss/retry cells the measured success rate is within {MAXDEV_SUCC:.1f} percentage points of the model (binomial standard error at most {MAXSE:.1f} points with n = {NQ:,} queries per cell). *Inference:* the success rate follows 1−(1−(1−p)²)^(r+1) closely (slightly below it at 30–40% loss), so the retry budget can be chosen analytically: two retries give ≥99% success up to 10% loss "
    f"({pc(dns_at(loss10, 2)['success'])} measured). Making the timeout adaptive (RFC 6298 style) keeps the same success rate but cuts the response time of lossy queries by "
    f"{min(RTO_SPEEDUP):.1f}–{max(RTO_SPEEDUP):.1f}× ({REF['rto']}).",
    f"**Reliability under failure.** All {T['total']} tests pass, including aborted uploads (no partial file remains), checksum corruption, path traversal, oversize inputs, malformed and looping DNS packets, idle-timeout and busy-server limits. "
    f"A crashed server is detected in {E7['crashed_server']['detected_after_s'] * 1000:.2f} ms, a hung one at the configured timeout ({f3(E7['hung_server']['detected_after_s'])} s for 1.0 s), and a dead UDP peer only after the retry budget ({f3(E7['udp_dead_server']['detected_after_s'])} s).",
    f"**Resources.** Idle server RSS is {f1(idle0['rss_mb'])} MB; each open connection adds ≈{f1(idle2000['kb_per_conn'])} KB and one thread; server CPU follows load (1% at 1 client, ≈{f0(by(E2, 'clients', 50)['cpu_pct'])}% at 50). "
    f"Memory under 400 active clients is {f1(c400['rss_peak_mb'])} MB.",
])
H2("b. Innovation, Novelty, and Evidence")
table(["Contribution", "What is new / non-trivial", "Evidence in this report"],
      [["Real DNS wire-format server in the same platform as echo and file services", "Hand-written RFC 1035 codec with name compression, CNAME chasing, NODATA/NXDOMAIN, TC handling, FORMERR/NOTIMP, loop-safe parser; forwarding resolver with TTL cache", f"Interoperates with `dig` ({REF['term_dns.png']}, test `matches_system_dig`); {sum('TestDNS' in t['id'] for t in T['tests'])} DNS tests"],
       ["Built-in, reproducible loss injection", "`--loss p` drops datagrams on receive and send, giving exact ground truth for reliability experiments", f"Measured loss within {max(abs(r['udp_observed_loss'] - r['udp_theory_loss']) for r in E5) * 100:.1f} pp of theory ({REF['loss_tab']}, {REF['fig_loss.png']})"],
       ["Closed-form retry model validated experimentally", "Success = 1−(1−(1−p)²)^(r+1) predicts DNS reliability and lets the retry budget be dimensioned", f"{REF['loss_tab']}: measured vs theory in every cell"],
       ["Atomic, verified file transfer", "SHA-256 announced before data, verified on both ends, temp file + atomic rename, cleanup on abort", "Tests on checksum mismatch, aborted upload, concurrent uploads"],
       ["Measured, root-caused optimisations", "Found a download/upload asymmetry through measurement, explained it (double file read), fixed it with a safe digest cache, and quantified `sendfile()`", f"+{(dl_dc100 / dl_base100 - 1) * 100:.0f}% download throughput; {c8[3]['cpu_pct'] / c8[0]['cpu_pct'] * 100:.0f}% CPU at 8 clients ({REF['file_conc']}, {REF['fig_files.png']})"],
       ["Adaptive retransmission timer", "RFC 6298 estimator (SRTT/RTTVAR, Karn’s rule, exponential back-off) applied to UDP request/response", f"Same success, {min(RTO_SPEEDUP):.1f}–{max(RTO_SPEEDUP):.1f}× lower latency ({REF['rto']}, {REF['fig_rto.png']})"],
       ["Honest scalability analysis", "Connection-burst limit, GIL ceiling, per-connection memory and thread cap all measured and explained", f"{REF['conc']}, {REF['idle']}; Section 9.e"],
       ["Observability", "`STATS` command returns live counters, CPU and memory; key=value logs are machine-parseable", f"{REF['term_concurrency.png']}, {REF['term_log.png']}"]],
      widths=[1.7, 2.7, 2.0], cap="Innovation and evidence summary", size=9)

# =============================================================================== 11 Conclusion
H1("11. Conclusion")
P("The project achieved its objectives. A single server process provides Echo (TCP and UDP), File Transfer and DNS services; serves hundreds of concurrent clients without errors; records every connection, request and "
  "failure in a structured log with live statistics; and detects and handles the failure modes that matter for a networked service, from a client that vanishes mid-upload to a malformed DNS packet. "
  "The evaluation produced quantitative answers on all five required metrics, and the loss experiments validate a simple analytical model, which turns retransmission policy into an engineering calculation instead of a guess.")
P("Two lessons stand out. First, *measurement exposes problems that reading code does not*: the upload/download asymmetry and the connection-burst failure were both found by benchmarks and both led to concrete changes. "
  "Second, *the choice between TCP and UDP is about who carries the reliability burden*: TCP hides loss at the cost of connection state and set-up time; UDP is cheaper per exchange but the application must "
  "supply timeouts and retries, and the data show exactly how many are needed.")

# =============================================================================== 12 Limitations
H1("12. Limitations and Future Enhancement")
H3("Limitations")
bullets([
    "Measurements are on the loopback interface only; real networks add latency, jitter, queueing, MTU effects and genuine loss.",
    "The thread-per-connection design is bounded by the interpreter’s single-core execution (≈29k echo requests/s) and by thread and memory limits (≈54 KB per connection; ≈4,000 threads per process on this macOS machine).",
    "A burst of more than roughly 200 simultaneous new connections can overflow the kernel accept queue (default `somaxconn` 128).",
    "No encryption or authentication: any client that can reach the ports can read, write and list files.",
    "The DNS service is authoritative only: no recursion (apart from simple forwarding), no EDNS0, no DNS over TCP for truncated answers, no DNSSEC, no wildcard records.",
    "The file service has no resume, quotas or directories. The adaptive UDP timeout uses a 5 ms floor tuned for a LAN, not the 1 s minimum RFC 6298 recommends for the open Internet.",
])
H3("Future enhancements")
bullets([
    "Run the same experiments across two hosts or with Linux `tc netem` / Mininet to add real delay and loss, and capture traffic with Wireshark to confirm the packet-level behaviour.",
    "Add TLS (`ssl`) and token-based authentication; per-client rate limiting.",
    "Raise `kern.ipc.somaxconn` and the listen backlog; replace thread-per-connection with an `asyncio` or `epoll` event loop, or several processes sharing a port with `SO_REUSEPORT`, and measure the effect.",
    "Per-destination RTT estimators shared across requests; DNS over TCP, EDNS0 and negative caching.",
    "Resumable, chunked file transfer with parallel streams; compare TCP congestion-control algorithms (Cubic, BBR) on the transfer service.",
    "Export metrics in Prometheus format and add a small dashboard.",
])

# =============================================================================== References
H1("References")
refs = [
    "J. Postel, “Transmission Control Protocol,” RFC 793, 1981; W. Eddy (ed.), “Transmission Control Protocol (TCP),” RFC 9293, 2022.",
    "J. Postel, “User Datagram Protocol,” RFC 768, 1980.",
    "L. Eggert, G. Fairhurst, G. Shepherd, “UDP Usage Guidelines,” RFC 8085, 2017.",
    "W. R. Stevens, B. Fenner, A. M. Rudoff, *UNIX Network Programming, Volume 1: The Sockets Networking API*, 3rd ed., Addison-Wesley, 2003.",
    "J. F. Kurose, K. W. Ross, *Computer Networking: A Top-Down Approach*, 8th ed., Pearson, 2021.",
    "J. Postel, “Echo Protocol,” RFC 862, 1983.",
    "J. Postel, J. Reynolds, “File Transfer Protocol (FTP),” RFC 959, 1985.",
    "K. Sollins, “The TFTP Protocol (Revision 2),” RFC 1350, 1992.",
    "P. Mockapetris, “Domain Names – Concepts and Facilities,” RFC 1034, 1987.",
    "P. Mockapetris, “Domain Names – Implementation and Specification,” RFC 1035, 1987.",
    "D. Kegel, “The C10K problem,” 1999. http://www.kegel.com/c10k.html",
    "V. Paxson, M. Allman, J. Chu, M. Sargent, “Computing TCP’s Retransmission Timer,” RFC 6298, 2011.",
    "Python Software Foundation, “socket – Low-level networking interface” and “threading”, Python 3 documentation. https://docs.python.org/3/library/socket.html",
    "R. Jain, *The Art of Computer Systems Performance Analysis*, Wiley, 1991.",
]
for i, r_ in enumerate(refs, 1):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.4)
    p.paragraph_format.first_line_indent = Inches(-0.4)
    p.paragraph_format.space_after = Pt(3)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    runs(p, f"[{i}]\t{r_}", size=11)

doc.core_properties.title = TITLE
doc.core_properties.author = ""
doc.core_properties.comments = ""
doc.core_properties.subject = "BCSE308P – Computer Networks, Project Based Component"
doc.save(OUT)
print("saved", OUT, f"({counters['fig']} figures, {counters['tab']} tables)")
